from pathlib import Path
import csv
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    '''
    Конвертирует CSV в XLSX с помощью openpyxl.
    Первая строка CSV - заголовок.
    Лист называется "Sheet1".
    Ширина колонок - автоширина по длине текста, но не менее 8 символов.
    '''
    csv_file = Path(csv_path)
    xlsx_file = Path(xlsx_path)

    if csv_file.suffix.lower() != ".csv":
        raise ValueError(f"Входной файл не .csv, а {csv_file.suffix.lower()}")
    if xlsx_file.suffix.lower() != ".xlsx":
        raise ValueError(f"Выходной файл не .xlsx, a {xlsx_file.suffix.lower()}")

    if not csv_file.exists():
        raise FileNotFoundError("Не найден входной CSV")

    with open(csv_file, "r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        data = list(reader)
        if not data:
            raise ValueError("CSV пуст")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    for row in data:
        ws.append(row)
    
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(8, max_len)
    
    xlsx_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_file)

if __name__=="__main__":
    csv_to_xlsx("data/lab05/samples/people.csv", "data/lab05/out/people_from_csv.xlsx")
    csv_to_xlsx("data/lab05/samples/cities.csv", "data/lab05/out/cities_from_csv.xlsx")

